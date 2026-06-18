"""
Генератор тестовых документов для DocScan.
Создаёт файлы со встроенными ПДн: паспорт, СНИЛС, ИНН, карты, телефоны, email.
"""

from pathlib import Path

OUT = Path(__file__).parent

# Реальные валидные тестовые номера (контрольные суммы рассчитаны правильно):
#   СНИЛС: 112-233-445 95 (валидный)
#   ИНН физ: 500100732259 (валидный, из тестов)
#   ИНН юр:  7736050003 (Газпром)
#   Карта:   4111 1111 1111 1111 (тестовая Visa, проходит Луна)


# ─── Word (DOCX) ──────────────────────────────────────────────────────────────

def make_docx():
    from docx import Document
    doc = Document()

    doc.add_heading("Заявление о приёме на работу", 0)

    doc.add_paragraph("В отдел кадров ООО «Тестовая компания»")
    doc.add_paragraph("от Иванова Ивана Ивановича")
    doc.add_paragraph("")

    doc.add_heading("Личные данные", 1)
    doc.add_paragraph("Фамилия, имя, отчество: Иванов Иван Иванович")
    doc.add_paragraph("Дата рождения: 15.03.1985 г.р.")
    doc.add_paragraph("Паспорт 4507 123456 выдан ОВД района 12.05.2010")
    doc.add_paragraph("СНИЛС: 112-233-445 95")
    doc.add_paragraph("ИНН налогоплательщика: 500100732259")
    doc.add_paragraph("Адрес регистрации: г. Москва, ул. Пушкина, д. 10, кв. 5")
    doc.add_paragraph("")

    doc.add_heading("Контактная информация", 1)
    doc.add_paragraph("Телефон мобильный: +7 (999) 123-45-67")
    doc.add_paragraph("Email: ivan.ivanov@example.com")
    doc.add_paragraph("")

    doc.add_heading("Платёжные реквизиты", 1)
    doc.add_paragraph("Номер карты для перечисления заработной платы: 4111 1111 1111 1111")
    doc.add_paragraph("Полис ОМС: 1234 5678 9012 3456")
    doc.add_paragraph("Водительское удостоверение: 50 12 345678")
    doc.add_paragraph("")

    doc.add_paragraph("Подпись: __________________ / И.И. Иванов /")
    doc.add_paragraph("Дата: 01.10.2026")

    path = OUT / "test_zayavlenie.docx"
    doc.save(str(path))
    print(f"OK: {path}")


# ─── Excel (XLSX) ─────────────────────────────────────────────────────────────

def make_xlsx():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Сотрудники"

    headers = [
        "№", "ФИО", "Дата рождения", "Паспорт", "СНИЛС", "ИНН",
        "Телефон", "Email", "Карта", "Адрес"
    ]
    ws.append(headers)

    employees = [
        [1, "Иванов Иван Иванович", "15.03.1985", "4507 123456", "112-233-445 95",
         "500100732259", "+7 (999) 123-45-67", "ivan@example.com",
         "4111 1111 1111 1111", "г. Москва, ул. Пушкина 10"],
        [2, "Петров Пётр Петрович", "22.07.1990", "4510 654321", "112-233-445 95",
         "500100732259", "+7 (916) 555-12-34", "petrov@company.ru",
         "5555 5555 5555 4444", "г. Санкт-Петербург, Невский 25"],
        [3, "Сидорова Анна Сергеевна", "05.11.1992", "4520 111222", "112-233-445 95",
         "500100732259", "+7 (903) 777-88-99", "anna.sidorova@firm.com",
         "4111 1111 1111 1111", "г. Казань, Кремлёвская 1"],
    ]

    for row in employees:
        ws.append(row)

    # Второй лист с банковскими данными
    ws2 = wb.create_sheet("Реквизиты")
    ws2.append(["Организация", "ИНН", "Контактное лицо", "Тел"])
    ws2.append(["ПАО Газпром", "7736050003", "Сергей Иванов", "+7 (495) 719-30-01"])
    ws2.append(["ООО Тест", "7743013902", "Мария Петрова", "+7 (812) 123-45-67"])

    path = OUT / "test_sotrudniki.xlsx"
    wb.save(str(path))
    print(f"OK: {path}")


# ─── RTF ──────────────────────────────────────────────────────────────────────

def make_rtf():
    # RTF — упрощённый формат, в DocScan читается как text/rtf
    rtf = r"""{\rtf1\ansi\ansicpg1251\deff0
{\fonttbl{\f0 Times New Roman;}}
\f0\fs24
{\b\fs32 Договор оказания услуг\par}
\par
Заказчик: Иванов Иван Иванович\par
Паспорт серии 4507 123456 выдан 12.05.2010\par
СНИЛС: 112-233-445 95\par
ИНН: 500100732259\par
Адрес: г. Москва, ул. Тверская, д. 1\par
Телефон: +7 (999) 123-45-67\par
Email: ivan@test.ru\par
\par
Реквизиты для оплаты:\par
Номер карты: 4111 1111 1111 1111\par
\par
Подпись: ___________ \par
}
"""
    path = OUT / "test_dogovor.rtf"
    path.write_text(rtf, encoding="cp1251")
    print(f"OK: {path}")


# ─── Изображение (PNG) с текстом — для теста OCR ──────────────────────────────

def make_image():
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        print("SKIP: Pillow не установлен")
        return

    img = Image.new("RGB", (800, 600), color="white")
    draw = ImageDraw.Draw(img)

    # Пытаемся загрузить системный шрифт
    font = None
    for fname in ["arial.ttf", "C:/Windows/Fonts/arial.ttf", "DejaVuSans.ttf"]:
        try:
            font = ImageFont.truetype(fname, 22)
            break
        except OSError:
            continue
    if font is None:
        font = ImageFont.load_default()

    lines = [
        "Анкета клиента",
        "",
        "ФИО: Иванов Иван Иванович",
        "Дата рождения: 15.03.1985 г.р.",
        "Паспорт: 4507 123456",
        "СНИЛС: 112-233-445 95",
        "ИНН: 500100732259",
        "",
        "Контакты:",
        "Телефон: +7 (999) 123-45-67",
        "Email: ivan@example.com",
        "",
        "Карта: 4111 1111 1111 1111",
    ]
    y = 30
    for line in lines:
        draw.text((40, y), line, fill="black", font=font)
        y += 35

    path = OUT / "test_anketa.png"
    img.save(str(path))
    print(f"OK: {path}")


if __name__ == "__main__":
    make_docx()
    make_xlsx()
    make_rtf()
    make_image()
    print(f"\nВсе файлы в: {OUT}")
